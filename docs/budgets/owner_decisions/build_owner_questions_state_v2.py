#!/usr/bin/env python3
"""Owner-question state v2 (A9-10, fo_a9_10_integration): deterministic builder.

Writes, next to this script:
    owner_questions_state_v2.json   machine-readable state of every owner question
    OWNER_QUESTIONS_STATE_v2.md     companion document generated from the JSON
    owner_questions_state_v2.csv    same rows, one per question
    owner_questions_state_v2.xlsx   same rows with a yellow 'Your answer' column on the OPEN rows only

Rows:
  * the 147 rows of the v1 consolidated list (owner_questions_consolidated.csv, byte-identical, never rewritten):
    ANSWERED (row pointer into docs/decisions/OD_2026_09_29_owner_answers_147.json) or SUPERSEDED (the owner's answer
    classifies the question as superseded for the primary campaign);
  * every A9.1 follow-up decision (docs/decisions/OD_2026_09_30_A9_1_followup_owner_decisions.json): ANSWERED (A9.1);
  * every question raised by the A9 lanes (A9-01..A9-09 open_owner_questions, the step-1 integration record, the M16
    v2 matrix and A9-10 itself), read from the current deliverables: ANSWERED_BY_A9_1 where the lane records that an
    A9.1 decision answers it, ANSWERED (row N) where an owner row answers it, ADDRESSED_IN_A9_10 for A9-10 work items,
    otherwise OPEN with the lane's proposed answer and needed-by gate. This builder never answers an OPEN question.

Usage:  python docs/budgets/owner_decisions/build_owner_questions_state_v2.py [--check]
The xlsx is compared by cell content in --check (its zip container carries timestamps).
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
REL = "docs/budgets/owner_decisions"
SCRIPT_REL = REL + "/build_owner_questions_state_v2.py"
JSON_REL, MD_REL = REL + "/owner_questions_state_v2.json", REL + "/OWNER_QUESTIONS_STATE_v2.md"
CSV_REL, XLSX_REL = REL + "/owner_questions_state_v2.csv", REL + "/owner_questions_state_v2.xlsx"
TEST_REL = "tests/test_owner_questions_state_v2.py"
V1_CSV = REL + "/owner_questions_consolidated.csv"
PINS = {
    "docs/decisions/OD_2026_09_29_owner_answers_147.json":
        "50e39a4deac7d4ada4710b2f641d717f1c4febd59366cbc04d8c66de6b4532b1",
    "docs/decisions/OD_2026_09_29_OWNER_DECISION_PACK_147.md":
        "8736d88a64bf26bd06a332a68c2450a780f175fffafe5aca0625759133666976",
    "docs/decisions/OD_2026_09_30_A9_1_followup_owner_decisions.json":
        "7a8f93dbc2487de90ebba0b2801fc5d3f5d983fc96ba418b55c492f1f9e851a4",
    "docs/decisions/OD_2026_09_30_A9_1_FOLLOWUP_OWNER_DECISIONS.md":
        "2587ca6931f6c9dac865005db9dc518dab0fcb829d789293467dc4179879c46e",
    "docs/decisions/OD_HARDWARE_PIVOT_2026_09_29_A9_hall_downstream_rf_icp_neutralizer.json":
        "74ef1a727c3656841ef115122c6d60865f7d2d93cfa29f7fb0081886484d2a1f",
    V1_CSV: "4bbd1cf4c97376c0e3f9f00353ee41759fe42956f9089767a39e669846fbee70",
    REL + "/OWNER_QUESTIONS_CONSOLIDATED.md": "a3d15db8bd332312a6db77a1fdd740b7e9f5972475f269fc43a76ddf7e1b062c",
    "docs/budgets/subsystem_maturity/subsystem_maturity_v2.json":
        "a82b1acd118b26e89edd4fa467bec778fa410470cca5eb6f684e6553eadaf23c",
}
ANS_REL = "docs/decisions/OD_2026_09_29_owner_answers_147.json"
A91_REL = "docs/decisions/OD_2026_09_30_A9_1_followup_owner_decisions.json"
M16_V2 = "docs/budgets/subsystem_maturity/subsystem_maturity_v2.json"
XE_JSON = "docs/budgets/" + "xe" + "_ledger_a9/" + "xe" + "_ledger_a9_v1.json"
LANE_SOURCES = [
    ("A9-01", "docs/experiments/hall_icp/prereg_framework/hall_icp_prereg_framework_v1.json"),
    ("A9-02", "docs/architecture_comparison/power_boundary_a9/bus_power_boundary_a9_v1.json"),
    ("A9-03", "schemas/interfaces/icp_neutralizer_icd_v1.json"),
    ("A9-04", "docs/experiments/hall_icp/uncertainty_budget/hall_icp_uncertainty_budget_v1.json"),
    ("A9-05 (evidence)", "docs/evidence/icp_neutralizer/icp_neutralizer_evidence_v1.json"),
    ("A9-05 (validation inputs)", "docs/experiments/hall_icp/validation_inputs/hall_icp_validation_inputs_v1.json"),
    ("A9-06", "docs/budgets/mass_a9/mass_a9_v1.json"),
    ("A9-07", "docs/hardware/h2_a9_revisions/h2_a9_revisions_v1.json"),
    ("A9-08", XE_JSON),
    ("A9-09", "docs/procurement/rfq_a9/rfq_a9_v1.json"),
    ("A9-INT (step-1 integration)", "docs/experiments/hall_icp/integration/a9_core_integration_v1.json"),
    ("A9-10", "docs/experiments/hall_icp/integration/a9_10_reconciliation_v1.json"),
    ("M16 v3 (A9-10)", "docs/experiments/hall_icp/integration/m16_v3/subsystem_maturity_v3.json"),
]
M16_ANSWERS = {"M16-Q-01": 140, "M16-Q-02": 141, "M16-Q-03": 142, "M16-Q-04": 143}
STATUSES = {
    "ANSWERED": "answered by the owner: a v1 row (row pointer into the 147 answers), an A9.1 decision row, or a lane "
                "question answered by an existing owner row (see status_detail)",
    "SUPERSEDED": "v1 row whose owner answer supersedes the question for the primary campaign",
    "ANSWERED_BY_A9_1": "lane question answered by an A9.1 decision",
    "ADDRESSED_IN_A9_10": "integration work item performed by A9-10 as directed by the brief (not an owner decision; "
                          "lane-assigned, the owner may reopen it)",
    "OPEN": "owner call; the lane's proposed answer is a proposal only"}
HEAD = ["No", "Id", "Source", "Question", "Status", "Answer / pointer", "Proposed (lane)", "Needed by", "Your answer"]


def _sha(rel: str) -> str:
    with open(os.path.join(ROOT, rel), "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def _load(rel: str):
    with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
        return json.load(f)


def verify_pins() -> list:
    out = []
    for rel, sha in PINS.items():
        got = _sha(rel)
        if got != sha:
            raise RuntimeError(f"pinned input changed: {rel} sha256 {got} != {sha}")
        out.append({"path": rel, "sha256": got})
    return out


def _short(s, n=240) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[: n - 3] + "..."


def v1_rows(answers: dict) -> list:
    with open(os.path.join(ROOT, V1_CSV), encoding="utf-8", newline="") as f:
        rdr = list(csv.DictReader(f))
    if len(rdr) != 147 or [int(r["No"]) for r in rdr] != list(range(1, 148)):
        raise RuntimeError("v1 consolidated list is not rows 1..147")
    out = []
    for r in rdr:
        n = int(r["No"])
        a = answers[n]
        if a["covers_ids"] != [x.strip() for x in r["Covers IDs"].split(",") if x.strip()]:
            raise RuntimeError(f"row {n}: covers_ids differ between the v1 list and the answers file")
        sup = a["classification_by_recorder"] == "SUPERSEDED"
        out.append({"no": n, "id": r["Covers IDs"], "source": f"v1 consolidated list row {n} ({r['Area']})",
                    "question": r["Question"], "status": "SUPERSEDED" if sup else "ANSWERED",
                    "answer_pointer": f"{ANS_REL} row {n}",
                    "answer_sha256": hashlib.sha256(a["owner_answer_verbatim"].encode("utf-8")).hexdigest(),
                    "answer_excerpt": _short(a["owner_answer_verbatim"]),
                    "classification_by_recorder": a["classification_by_recorder"],
                    "proposed": r["Proposed"], "needed_by": r["Needed by"], "kind": "v1_row"})
    return out


def a91_rows(a91: dict, lane_q: list) -> list:
    out = []
    for did, text in a91["decisions"].items():
        base = did.replace("_accounting", "")
        answered = [q for q in lane_q if re.search(r"\b" + re.escape(base) + r"\b", q.get("_answered_by", ""))]
        rows = sorted({int(m) for q in answered for m in re.findall(r"\brows? (\d+)", q["question"])})
        out.append({"id": did, "source": "A9.1 follow-up owner decisions (" + A91_REL + ")",
                    "question": "A9.1 decision " + did, "status": "ANSWERED (A9.1)",
                    "answer_pointer": f"{A91_REL} decisions.{did}",
                    "answer_excerpt": _short(text if isinstance(text, str) else json.dumps(text)),
                    "answers_lane_questions": sorted(q["id"] for q in answered),
                    "refines_v1_rows_cited_by_those_questions": rows, "proposed": "", "needed_by": "",
                    "kind": "a9_1"})
    return out


def lane_rows(answers_file: dict) -> list:
    out = []
    for lane, rel in LANE_SOURCES:
        doc = _load(rel)
        for q in doc.get("open_owner_questions", []):
            st = str(q.get("status", ""))
            prop = q.get("proposed_answer") or q.get("proposed") or ""
            need = q.get("needed_by") or q.get("freeze_point")
            if not need:   # row 144: every blocker needs a stated latest decision point; derived, marked PROPOSED
                txt = q["question"] + " " + str(q.get("proposed_answer") or q.get("proposed") or "")
                gate = "LOCK-2" if "LOCK-2" in txt and "LOCK-1" not in txt else "LOCK-1"
                need = f"{gate} (PROPOSED by A9-10 per row 144; the lane stated no gate)"
            row = {"id": q["id"], "source": f"{lane} ({rel})", "question": q["question"], "proposed": prop,
                   "needed_by": need, "kind": "lane", "lane": lane, "_answered_by": ""}
            m91 = re.match(r"ANSWERED_BY_A9_1 \(([^)]*)\)", st)
            mrow = re.match(r"ANSWERED_BY_OWNER_ROW_(\d+)", st)
            if m91:
                row.update(status="ANSWERED_BY_A9_1", answer_pointer=f"{A91_REL} decisions.{m91.group(1)}",
                           answer_excerpt=_short(q.get("a9_1_decision")), _answered_by=m91.group(1))
            elif mrow:
                n = int(mrow.group(1))
                row.update(status=f"ANSWERED (owner row {n})", answer_pointer=f"{ANS_REL} row {n}",
                           answer_excerpt=_short(answers_file[n]["owner_answer_verbatim"]))
            elif st.startswith("ADDRESSED_IN_A9_10"):
                row.update(status="ADDRESSED_IN_A9_10", answer_pointer=_short(st, 300), answer_excerpt="")
            else:
                row.update(status="OPEN", answer_pointer="", answer_excerpt="")
            out.append(row)
    m16 = _load(M16_V2)
    for q in m16["open_owner_questions"]:
        n = M16_ANSWERS[q["id"]]
        if q["id"] not in answers_file[n]["covers_ids"]:
            raise RuntimeError(f"{q['id']} not covered by owner row {n}")
        out.append({"id": q["id"], "source": f"M16 v2 ({M16_V2})", "question": q["question"],
                    "status": f"ANSWERED (owner row {n})", "answer_pointer": f"{ANS_REL} row {n}",
                    "answer_excerpt": _short(answers_file[n]["owner_answer_verbatim"]), "proposed": "",
                    "needed_by": "", "kind": "lane", "lane": "M16 v2", "_answered_by": ""})
    return out


def build() -> dict:
    pins = verify_pins()
    answers = {a["row"]: a for a in _load(ANS_REL)["answers"]}
    a91 = _load(A91_REL)
    v1 = v1_rows(answers)
    lanes = lane_rows(answers)
    a9 = a91_rows(a91, lanes)
    rows = v1 + a9 + lanes
    for i, r in enumerate(rows, 1):
        r["no"] = i
        r.pop("_answered_by", None)
        detail = r["status"]
        r["status"] = detail.split(" (")[0]        # canonical, machine-filterable status (review repair)
        r["status_detail"] = detail                 # e.g. 'ANSWERED (A9.1)', 'ANSWERED (owner row 111)'
        if r["status"] not in STATUSES:
            raise RuntimeError(f"{r['id']}: status {detail!r} outside the vocabulary")
    ids = [(r["id"], r["source"]) for r in rows]
    if len(ids) != len(set(ids)):
        raise RuntimeError("duplicate question row")
    counts = {}
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    open_rows = [r for r in rows if r["status"] == "OPEN"]
    return {
        "schema": "owner_questions_state_v2", "id": "owner_questions_state_v2", "lane": "fo_a9_10_integration",
        "trigger": "T_A9_10_INTEGRATION", "generated_by": SCRIPT_REL, "companion_document": MD_REL,
        "exports": [CSV_REL, XLSX_REL], "test": TEST_REL,
        "status": "STATE_RECORD_FOR_OWNER (answers come only from the owner)",
        "a9_status": "OWNER_AUTHORIZED_INVESTIGATION_HYPOTHESIS_NOT_FLIGHT_BASELINE",
        "supersedes_for_use": {"path": V1_CSV, "note": "v1 consolidated list (and its .md / .xlsx) kept byte-identical"},
        "pins": pins,
        "status_vocabulary": {k: v for k, v in STATUSES.items()},
        "status_detail_vocabulary": {
            "ANSWERED": "v1 row answered by the owner (row pointer into the 147 answers)",
            "ANSWERED (A9.1)": "A9.1 follow-up decision (binding); status ANSWERED",
            "ANSWERED (owner row N)": "lane question answered by an existing owner row; status ANSWERED",
            "SUPERSEDED / ANSWERED_BY_A9_1 / ADDRESSED_IN_A9_10 / OPEN": "detail equals the status"},
        "status_rule": "'status' is the canonical value (filter on it; 'counts' counts it); 'status_detail' keeps the "
                       "provenance form. ADDRESSED_IN_A9_10 is assigned by the lane for integration work the brief "
                       "directed A9-10 to perform (OQ-INT-03 / OQ-INT-04); it is not an owner decision and the owner "
                       "may reopen it.",
        "counts": dict(sorted(counts.items())), "open_count": len(open_rows),
        "rows": rows,
        "compliance": ["no OPEN question is answered here", "immutable inputs pinned by sha256; the v1 list is never "
                       "rewritten", "no winner declared"],
    }


# ------------------------------------------------------------------------------------------------------------------
def table(doc) -> list:
    return [[r["no"], r["id"], r["source"], r["question"], r["status"],
             (r.get("answer_pointer") or "") + ((" - " + r["answer_excerpt"]) if r.get("answer_excerpt") else ""),
             r.get("proposed") or "", r.get("needed_by") or "", ""] for r in doc["rows"]]


def render_md(doc) -> str:
    def esc(s):
        return str(s).replace("|", "\\|").replace("\n", " ")
    L = ["# Owner-question state v2 (A9-10)", "",
         f"<!-- GENERATED by {SCRIPT_REL}; do not edit by hand -->", "",
         f"{len(doc['rows'])} rows: the 147 v1 rows, every A9.1 decision and every question raised by the A9 lanes. "
         f"Counts: {', '.join(f'{k} {v}' for k, v in doc['counts'].items())}. **{doc['open_count']} OPEN** (owner calls; "
         "fill the yellow 'Your answer' column of the xlsx). No OPEN question is answered here.", "",
         "Status vocabulary: " + "; ".join(f"**{k}** = {v}" for k, v in doc["status_vocabulary"].items()) + ".", "",
         "## OPEN questions", "", "| No | Id | Source | Question | Proposed (lane) | Needed by |", "|---|---|---|---|---|---|"]
    for r in doc["rows"]:
        if r["status"] == "OPEN":
            L.append(f"| {r['no']} | {esc(r['id'])} | {esc(r['source'])} | {esc(r['question'])} | "
                     f"{esc(r['proposed'])} | {esc(r['needed_by'])} |")
    L += ["", "## All rows", "", "| " + " | ".join(HEAD[:-1]) + " |", "|" + "---|" * (len(HEAD) - 1)]
    for t in table(doc):
        L.append("| " + " | ".join(esc(c) for c in t[:-1]) + " |")
    L += ["", "Pins: " + "; ".join(f"`{p['path']}` `{p['sha256']}`" for p in doc["pins"])]
    return "\n".join(L) + "\n"


def render_csv(doc) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(HEAD)
    w.writerows(table(doc))
    return buf.getvalue()


def write_xlsx(doc, path) -> None:
    import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook()
    fixed = datetime.datetime(2026, 9, 30, 0, 0, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = "build_owner_questions_state_v2.py"
    ws = wb.active
    ws.title = "Owner question state v2"
    f = Font(name="Arial", size=10)
    fb = Font(name="Arial", size=10, bold=True)
    ws.append(["Owner-question state v2 - fill in the yellow 'Your answer' column (column I) of the OPEN rows only. "
               "'Proposed (lane)' is the lane's proposal, not a decision."])
    ws["A1"].font = fb
    ws.append(HEAD)
    for r in table(doc):
        ws.append(r)
    for i, w in enumerate([5, 18, 30, 70, 20, 60, 45, 18, 30]):
        ws.column_dimensions[chr(65 + i)].width = w
    yellow = PatternFill("solid", start_color="FFFF00")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = fb if c.row == 2 else f
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[0].row > 2 and row[4].value == "OPEN":
            row[8].fill = yellow
    ws.freeze_panes = "E3"
    t = Table(displayName="OwnerQuestionsV2", ref=f"A2:I{ws.max_row}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    wb.save(path)


def xlsx_cells(path) -> list:
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    out = []
    for row in ws.iter_rows():
        out.append([(c.value, c.fill.start_color.rgb if c.fill and c.fill.fill_type else None) for c in row])
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="verify the outputs reproduce")
    a = ap.parse_args(argv)
    doc = build()
    outs = {JSON_REL: json.dumps(doc, indent=1, ensure_ascii=False) + "\n", MD_REL: render_md(doc),
            CSV_REL: render_csv(doc)}
    if a.check:
        bad = [rel for rel, txt in outs.items()
               if not os.path.isfile(os.path.join(ROOT, rel))
               or open(os.path.join(ROOT, rel), encoding="utf-8", newline="").read() != txt]
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            tmp = os.path.join(td, "x.xlsx")
            write_xlsx(doc, tmp)
            if not os.path.isfile(os.path.join(ROOT, XLSX_REL)) or \
                    xlsx_cells(tmp) != xlsx_cells(os.path.join(ROOT, XLSX_REL)):
                bad.append(XLSX_REL)
        print("OK" if not bad else "DRIFT: " + ", ".join(bad))
        return 0 if not bad else 1
    for rel, txt in outs.items():
        with open(os.path.join(ROOT, rel), "w", encoding="utf-8", newline="") as f:
            f.write(txt)
    write_xlsx(doc, os.path.join(ROOT, XLSX_REL))
    print(f"wrote {JSON_REL}, {MD_REL}, {CSV_REL}, {XLSX_REL} ({len(doc['rows'])} rows, {doc['open_count']} OPEN)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
