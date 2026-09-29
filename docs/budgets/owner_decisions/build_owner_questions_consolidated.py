"""Consolidated owner-question list (every open question that needs an owner answer, one row per decision).

Sources (read 2026-09-29, execution branch): owner_decision_register_v1.json (44 items), LOCK-1 decision brief
(D-01..D-15), hardware_requirements_v1.json (HWQ-01..21), the seven H2 lane JSONs (docs/hardware/h2/), mass_bom_v1
(OD-M1..7), hard-gate OD1, web-track threads R3-R8 and EXTERNAL_EVIDENCE_PACKAGE.md sec. 5, the v2 RF||Hall branch
(feature/rf-hall-parallel-v2) and repository housekeeping. Duplicates across sources are merged into one row whose
'Covers IDs' column lists every source id it answers. Wording is condensed; the source files hold the full text.
Ids of the form H22-FQ-n, H23-Q-n, H25-Q6, H26Q-08, H27-Q7, REFRESH-Qn, WEB-*, V2-Qn, REPO-n are labels given here to
questions their source states without an id. This file decides nothing: 'Proposed' is what the source proposes (or 'none - owner call').

    python docs/budgets/owner_decisions/build_owner_questions_consolidated.py   # writes .md, .csv, .xlsx
"""
from __future__ import annotations

import csv
import os

HERE = os.path.dirname(os.path.abspath(__file__))
L1 = "LOCK-1"
N = "none - owner call"

# (area, covers_ids, question, options, proposed, needed_by, affects)
ROWS = [
    # ---- A. RFP, access, repository -------------------------------------------------------------------------------
    ("A. RFP / access / repo", "WEB-ACC-1", "Obtain the RFP document (defproc / tdf.drdo.gov.in need a human for the captcha / are 503).",
     "you download it / not", "you obtain it", "ASAP (bid close)", "every RFP reading below"),
    ("A. RFP / access / repo", "WEB-ACC-2", "Confirm RFP number (XI/LM vs X/L/M) and bid close date (05 Oct vs 20 Oct 2026).",
     "", N, "ASAP", "schedule"),
    ("A. RFP / access / repo", "WEB-RFP-1", "Mission life: 26,000 h (repo) vs 'three years' = 26,280 h (news); and is '> 15,000 h firing' in the RFP?",
     "", N, "ASAP", "life, Xe ledger (OD-XE-5)"),
    ("A. RFP / access / repo", "OD1", "Meaning of '12-25 mN'.",
     "(i) operating window, 25 mN cap | (ii) 12 mN sustained + 25 mN peak capability with Xe", N, L1, "hard gates, D-15, HWQ-10"),
    ("A. RFP / access / repo", "OD-XE-8, OD-M7, H27-Q3", "Does the 40 kg include the Xe load (and tank)?",
     "includes Xe + tank | dry mass only", "includes (verify against RFP)", "—", "mass closure, A7 blocker 3"),
    ("A. RFP / access / repo", "OD-XE-6", "Does 'air + Xe' / 'supplemented with Xenon' require Xe-mode operation beyond ignition, cathode and contingency?",
     "no | yes (then a new explicit Xe ledger term)", N, "—", "Xe mass, A7 blocker 3"),
    ("A. RFP / access / repo", "WEB-ACC-3, OQ-R8-3, OQ-R3-5", "Authorize legitimate acquisition (library / ILL / purchase / owner channel) of: Li 2015 (compressor CR), Tejeda & Knoll 2023, Bayliss & Knoll IEPC-2025, Cifali 2012, Wang 1995, blocked datasheets (Combat BN, C-1 candidates).",
     "yes (list) / no", N, "before Milestone B", "compressor C1 evidence, anode, materials"),
    ("A. RFP / access / repo", "WEB-SUP-1, R6-Q5", "Request supplier quotations at all (Xe tank 5-20 L, low-flow PMU/FCU, MFCs, thrust stand)? This means contacting suppliers.",
     "yes / no / which", N, "H3 wave", "procurement"),
    ("A. RFP / access / repo", "REPO-1", "Merge the H2 hardware wave (H2-1..H2-7, all verified; on claude/nifty-ramanujan-w68f9z) into main via a new PR?",
     "yes / not yet", N, "—", "repo (checkpoint approval was one-time)"),
    ("A. RFP / access / repo", "REPO-2", "Delete 8 leftover worktrees from failed runs (uncommitted partial edits, all superseded)?",
     "delete / keep", "delete", "—", "housekeeping"),
    ("A. RFP / access / repo", "REPO-3", "Parallel RF||Hall v2 branch (feature/rf-hall-parallel-v2, 7335dcd, reviewed): merge, open a PR, or keep as a branch?",
     "merge / PR / keep", N, "—", "repo"),

    # ---- B. LOCK-1 decisions (experiment design) --------------------------------------------------------------------
    ("B. LOCK-1 (experiment)", "D-01", "Split of the ln R_arch uncertainty budget over variance groups G1-G5.",
     "A equal five-way | B G5 removed, 4 equal (only with D-06-B) | C G3-G5 at 0.1, G1/G2 0.35", "D-01-A", L1, "instrument/mount specs"),
    ("B. LOCK-1 (experiment)", "D-02", "Arm stop rule D3.", "A sign form (upper bound < 0) | B STOP-MARGIN (< -delta)", "D-02-A", L1, "A7 blocker 2"),
    ("B. LOCK-1 (experiment)", "D-03", "Must a stop from the confirmation subset rest on Hall-on data at P_lo too?", "A no | B yes (+6 rows)", "D-03-B", L1, "campaign size"),
    ("B. LOCK-1 (experiment)", "D-04", "Iso-power chord error u_interp.", "A pre-registered Type B bound | B measure with V_mid reading (+2 conditions/block)", "D-04-B", L1, "uncertainty"),
    ("B. LOCK-1 (experiment)", "D-05", "Where the signed LOCK-1/LOCK-2 files live.",
     "A experiment_protocol/prereg | B minimum_decisive_experiment/prereg | C experiment_package/prereg | D lock1/", "D-05-A", L1, "governance"),
    ("B. LOCK-1 (experiment)", "D-06, HWQ-01", "How HW-0 / HW-RF / HW-ECR are realized on H-1; and configuration change = swap module only (H-1 stays on stand) or re-mount H-1 each time.",
     "A HW-0 spacer + S1b re-mounts | B in-vacuum diverter | C both applicators installed (CFG-A) | D separate builds with shams (CFG-B)", "D-06-A; module-only swap (HW-SVC-04)", L1, "A7 blocker 2"),
    ("B. LOCK-1 (experiment)", "D-07, R-08", "Decision margin delta on ln R_arch, and whether effect size is a LOCK-1 rule or a LOCK-2 value (A6 forbids numeric Phase-1 thresholds now).",
     "A delta 0.05 | B delta 0.1", "D-07-A (lane-25 T-DELTA)", L1, "A7 blocker 2"),
    ("B. LOCK-1 (experiment)", "D-08, R-03", "Admissible block counts n (n computed at LOCK-2); must n be whole replicates of the 6 sequences?",
     "A even n in [4,8] | B replicates_min = 3", "D-08-A; whole replicates or a registered incomplete-design rule", L1, "campaign size"),
    ("B. LOCK-1 (experiment)", "D-09, HWQ-02", "Tested Hall unit H-1: Vyovrinda design or a surrogate; definition of 'design-representative'.",
     "A Vyovrinda (level 1) | B surrogate (level 3) | C Vyovrinda + non-scoring surrogate pilot", "D-09-A", L1, "evidence level"),
    ("B. LOCK-1 (experiment)", "D-10", "Adopt a single protocol basis (lane 25 + lane 06 reconciliation)?", "A adopt | B run separately", "D-10-A", L1, "admissibility"),
    ("B. LOCK-1 (experiment)", "D-11", "Unmeasured common loads (compressor bus draw, valve-outlet feed state).",
     "A report PARTIAL_BOUNDARY until the ICD supplies them | B fix a compressor ledger input now", "D-11-A", L1, "bus power"),
    ("B. LOCK-1 (experiment)", "D-12, H26Q-01, R5-Q3", "Facility requirements per stage, T-PB-MAX (max background pressure), number of elevated p_b levels; fallback if ~1e-5 Torr at 3.2 mg/s (~212,000 L/s N2) is unreachable - accept 5e-5 Torr (~42,000 L/s) with the S5 slope reported?",
     "A one elevated level | B two", "D-12-A; T-PB-MAX: owner call (set with the low-flow knee in view)", L1, "facility choice, A7 blocker 1"),
    ("B. LOCK-1 (experiment)", "D-13", "Measure ignition / start attempts per arm (Xe-assisted and air-only)?", "A exclude | B add, non-score-bearing", "D-13-B", L1, "G6.ignition"),
    ("B. LOCK-1 (experiment)", "D-14", "Physics track pre-registers part of the new data as held-out Hall-transport validation?", "A no | B yes, own prereg before S1", "D-14-B", L1, "Hall validation"),
    ("B. LOCK-1 (experiment)", "D-15, R-14", "Scope extensions: Xe health check; Xe-augmented peak points (labelled XE_AUGMENTED_PEAK, never Case A evidence, Xe booked to the Xe ledger).",
     "A none | B Xe health check | C B + Xe-augmented peak points (only under OD1 (ii))", "D-15-C", L1, "Xe mass, 25 mN demo"),
    ("B. LOCK-1 (experiment)", "HWQ-10", "Absolute gate: must 25 mN capability be shown within P_bus < 1.5 kW, or only sustained 12 mN?",
     "25 mN (16.67 mN/kW floor) | 12 mN (8.0 mN/kW)", N, L1, "absolute thrust gate"),

    # ---- C. Phase-1 prereg framework --------------------------------------------------------------------------------
    ("C. Phase-1 framework", "P1F-OOD-01", "Accept the Phase-1 decision topology (cases A hall_only / B rf_hall / C ecr_hall / NO_VIABLE_CASE via P1DQ decision quantities)?",
     "accept | amend by dated addendum before S1 data", "accept", L1, "A7 blockers 1, 2"),
    ("C. Phase-1 framework", "R-01", "Configuration order: order-balanced design governs; old HW-0-first/last become reference checks?", "", "accept", "—", "measurement"),
    ("C. Phase-1 framework", "R-02", "Randomization seed timing.", "", "draw at LOCK-2 after n fixed; optional sha256 commitment at LOCK-1", L1, "measurement"),
    ("C. Phase-1 framework", "R-04", "Condition grid fixed at LOCK-1 from W1 + anticipated knee; knee is an output per configuration?", "", "accept", L1, "measurement"),
    ("C. Phase-1 framework", "R-05", "eta_u has no S1b repeatability.", "add Faraday/ExB repeatability at S1b | report eta_u descriptively (say how Case A is judged)", N, "—", "A7 blocker 1"),
    ("C. Phase-1 framework", "R-06", "Add RR-07 module-exchange remount check (RF/ECR install reproducibility)?", "", "yes", "LOCK-2", "A7 blocker 2"),
    ("C. Phase-1 framework", "R-07", "Confirm mapping: A5 'Phase 1' spans pivot Phase-1 knee, Phase-2 comparison and Phase-3 absolute gates.", "", "you confirm", L1, "A7 blockers 1, 2"),
    ("C. Phase-1 framework", "R-09", "W5 class of HW0-REF readings.", "", "held-out custody until W5 decides", "—", "none"),
    ("C. Phase-1 framework", "R-10", "Within-installation gas order.", "", "pure N2 first, O2-bearing last; Xe reference per D-15", "—", "measurement"),
    ("C. Phase-1 framework", "R-11", "Form of the Case B net system-level benefit where both sustain.", "", "you fix NET_BENEFIT form at LOCK-1", L1, "A7 blocker 2"),
    ("C. Phase-1 framework", "R-12, P1F-OOD-05", "Evidence that establishes no outcome.", "", "decision status OPEN (not an outcome)", L1, "A7 blockers 1, 2"),
    ("C. Phase-1 framework", "R-13", "An arm stop breaks schedule balance.", "", "skip stopped arm's slots (NOT_TESTED), keep others' order, report balance loss", "—", "measurement"),
    ("C. Phase-1 framework", "P1F-OOD-03", "Reference condition REF-COND (S1b re-mount point: nominal flow/voltage, N2): own installation per block, or merged into adjacent HW-0 (REF-MERGED)?",
     "own installation | REF-MERGED", "own installation (default)", L1, "measurement"),
    ("C. Phase-1 framework", "P1F-OOD-04", "MD-04: abort on a thruster/configuration limit = 'not sustained within limits' (reported with the limit)?", "", "yes", L1, "A7 blocker 1"),

    # ---- D. Xe system ----------------------------------------------------------------------------------------------
    ("D. Xe system", "OD-XE-1, R6-Q2", "Xe accounting convention; is Xe flowing during cathode preheat booked?",
     "PHASE_TOTAL_FLOW | CATHODE_CONTINUOUS_INCREMENTAL", "PHASE_TOTAL_FLOW (conservative)", "—", "Xe mass, A7 blocker 3"),
    ("D. Xe system", "OD-XE-2", "Reserve policy (form and value).", "fraction of other terms | absolute mass", "own term; value: owner call", "—", "Xe mass"),
    ("D. Xe system", "OD-XE-3", "Share of the mass reference the stored-Xe subsystem may take.", "0.25 / 0.5 / 1.0", "0.25 screening", "—", "Xe mass"),
    ("D. Xe system", "OD-XE-4, H27-Q6", "2 % residual (unusable) Xe: separate line in the Xe ledger, keep in mass BOM only, or fold into reserve?",
     "Xe ledger line | mass BOM only | fold into m_reserve", N + " (never booked twice)", "—", "Xe mass"),
    ("D. Xe system", "OD-XE-5", "Firing-hours basis of the cathode term.", "15,000 h (5.4 kg) | 18,000 h (6.48 kg)", "15,000 h (A5)", "—", "Xe mass"),
    ("D. Xe system", "OD-XE-7", "Xe hardware-to-mass-BOM mapping (regulator/valves/plumbing -> xe_valve_and_flow_control, tank -> xe_tank).", "", "accept", "—", "mass"),
    ("D. Xe system", "OD-M4", "Declare a mission Xe load (design input).", "", N, "—", "tank sizing"),
    ("D. Xe system", "R6-Q1", "C-1 heated or heaterless (heaterless per-start Xe 2-6x higher)?", "heated | heaterless", N, "—", "Xe mass, PPU"),
    ("D. Xe system", "R6-Q3", "Tank sizing at a stated maximum storage temperature (1.67 g/cm3 at 323 K / 150 bar vs 1.97 at 300 K)?", "", N, "—", "tank volume"),
    ("D. Xe system", "R6-Q4", "Filter-getter (<= 17 W) in the Xe cathode line in scope?", "yes | no", N, "—", "power, mass"),

    # ---- E. Mass ---------------------------------------------------------------------------------------------------
    ("E. Mass", "OD-M1, H27-Q2", "System mass margin.", "20 % (ESA analogy) | 10 % (RFP R3, verify)", "20 %", "—", "mass closure"),
    ("E. Mass", "H27-Q1", "Evaluate the allocation at both ends of '<= 34-36 kg', or fix one value?", "both | 34 | 36", "both (as done)", "—", "mass"),
    ("E. Mass", "H27-Q7", "Set per-row mass allocations now (e.g. PPU, regulator) so H2-3/H2-4 size against a budget? (High analog case exceeds 40 kg on known lines.)", "yes (values) | no", N, "—", "mass closure"),
    ("E. Mass", "H27-Q4", "Redundancy policy for valves, cathode and PPU.", "single string | redundant (which)", "single string (assumed)", "—", "mass"),
    ("E. Mass", "H27-Q5", "Accept NASA SBIR 2 kg PPU development target (undemonstrated) as PPU low end until H2-4 has a CBE?", "", N, "—", "mass"),
    ("E. Mass", "OD-M2", "Maturity category for dry equipment.", "ECSS D (20 %) for all | per selected part", "D for all", "—", "mass"),
    ("E. Mass", "OD-M3", "Legacy uncited MGA dict in abep_sim/mass_bom.py.", "retire (model change, goldens move) | keep and cite", N, "—", "model"),
    ("E. Mass", "OD-M5", "Mass BOM item list.", "accept | amend", "accept", "—", "mass"),
    ("E. Mass", "OD-M6", "Harness allocation.", "5 % of dry mass (ESA) | designed harness", "5 %", "—", "mass"),

    # ---- F. Pre-ionizer ICD and module hardware --------------------------------------------------------------------
    ("F. Pre-ionizer ICD", "PMQ-01", "Size the common module envelope to the largest occupant (ECR with magnet/yoke/feed)?", "", "approve", "—", "A7 blocker 2"),
    ("F. Pre-ionizer ICD", "PMQ-02", "Common control harness L1-L5 incl. MODULE_ID and blank-module termination?", "", "approve", "—", "A7 blocker 2"),
    ("F. Pre-ionizer ICD", "PMQ-03", "Pressure-drop class definition, matched passive inserts for the blank module if needed?", "", "approve", "—", "A7 blockers 1, 2"),
    ("F. Pre-ionizer ICD", "PMQ-04", "Add S1a no-plasma module-exchange series for RF/ECR modules (cold flow, zero/tare, B(z))?", "", "approve", "—", "A7 blocker 2"),
    ("F. Pre-ionizer ICD", "PMQ-05", "Time-matched Xe hold in the HW-0 start equal to the pre-ionizer dwell?", "", "approve", "—", "A7 blocker 2"),
    ("F. Pre-ionizer ICD", "PMQ-06", "Module loads with no bus-power slot (RF assist magnet, powered interstage, coolant pump).", "prohibit in v1 | new bus-boundary version", N, L1, "bus power"),
    ("F. Pre-ionizer ICD", "HWQ-04", "INV-B3 tolerance on module-induced field change; add coil-current sensitivity scan S_B to LOCK-1?", "", N, L1, "A7 blocker 2"),
    ("F. Pre-ionizer ICD", "HWQ-05", "ECR magnet type.", "electromagnet (coil power booked) | permanent magnet (0 W)", N + " (must match flight intent)", L1, "bus power"),
    ("F. Pre-ionizer ICD", "HWQ-06", "RF module magnetization.", "unmagnetized ICP in v1 | define where assist-magnet power is booked", "unmagnetized ICP", L1, "bus power"),
    ("F. Pre-ionizer ICD", "HWQ-07", "Electrical potential of module body / interstage electrode.", "floating | anode | cathode common | facility ground", N, L1, "A7 blocker 2"),
    ("F. Pre-ionizer ICD", "HWQ-15", "Freeze module interfaces before Phase 1 even if modules arrive later?", "", "yes", L1, "A7 blocker 2"),
    ("F. Pre-ionizer ICD", "R4-Q7", "RF/ECR: select frequency and P_hi (to specify INS-03 sensors); accept calorimetric load-plane method if inline sensors can't reach u_src?", "", N, "—", "instrumentation"),

    # ---- G. H-1 Hall head, magnet, thermal (H2-1, H2-5) ------------------------------------------------------------
    ("G. H-1 head / magnet / thermal", "H2-1 Q1", "Design flow for channel sizing (delivered flow 0.38-1.29 mg/s design case up to ~3.2 mg/s): where to match channel neutral density?", "", N, "H2 freeze", "H-1 geometry"),
    ("G. H-1 head / magnet / thermal", "H2-1 Q2", "Keep magnetically shielded topology (T2) as H-1 baseline, with a replaceable unshielded (T1) pole set?", "", "yes", "H2 freeze", "H-1, life"),
    ("G. H-1 head / magnet / thermal", "H2-1 Q3", "L/h upper bound (12 proposed); allow an adjustable-anode development insert before S1?", "", "12; owner call on insert", "H2 freeze", "H-1 geometry"),
    ("G. H-1 head / magnet / thermal", "HWQ-18, H2-1 Q4, H25-Q5, OQ-R3-1", "Pole/core material grade (e.g. Hiperco 50 / FeCo-2V inner, pure iron outer, or 430F) and a max-use pole temperature below Curie set from B_sat(T).",
     "pure iron | FeCo-2V | 430F SQ", "FeCo-2V inner + pure iron outer; temperature limit adopted", "S1a", "magnet, thermal"),
    ("G. H-1 head / magnet / thermal", "HWQ-20, H2-1 Q5, H25-Q2, OQ-R3-2, REFRESH-Q2", "Coil insulation family / IEC 60085 class (ceramic Ni-clad Cu vs polyimide), hot-spot margin; accept ~150 V AC turn rating and Ni ferromagnetism if ceramic; decide before or at LOCK-1?",
     "ceramic Ni-clad Cu | polyimide", N, "S1a (MCQ-S1-01)", "magnet (blocks M16 magnetic_circuit)"),
    ("G. H-1 head / magnet / thermal", "HWQ-19", "MC-1: electromagnet only (traceable B(z) vs current) or permanent-magnet assisted?", "EM only | PM assisted", N, "—", "magnet, validation"),
    ("G. H-1 head / magnet / thermal", "H2-1 Q6, H22-OQ-01", "Cathode location: central (L-CENTRAL; constrains mean diameter) or external (L-EXTERNAL)?", "central | external", "central (preliminary)", "H2 freeze", "H-1, C-1"),
    ("G. H-1 head / magnet / thermal", "H2-1 Q7", "Proposed design factors: B headroom 1.5, working flux 0.7 B_sat, discharge share >= 0.5.", "", "accept", "H2 freeze", "magnet sizing"),
    ("G. H-1 head / magnet / thermal", "HWQ-03", "Rate H-1, C-1, supply and isolator to the relaxed 350 V V_d end?", "", "yes", "—", "hardware rating"),
    ("G. H-1 head / magnet / thermal", "HWQ-13", "Hot-state B reference sensor on MC-1?", "", "yes", "—", "validation"),
    ("G. H-1 head / magnet / thermal", "HWQ-14", "Spares policy: replaced/repaired H-1 or C-1 = new unit with new HW-0 reference and S1b?", "", "yes", "—", "test validity"),
    ("G. H-1 head / magnet / thermal", "H25-Q1", "Accept a high-emittance temperature-capable exterior coating on MC-1 as a baseline requirement (bare metal is design-driving)?", "", "yes", "—", "thermal"),
    ("G. H-1 head / magnet / thermal", "H25-Q3", "Spacecraft thermal ICD: mounting-interface temperature and allowable heat into the spacecraft.", "", N, "—", "thermal"),
    ("G. H-1 head / magnet / thermal", "H25-Q4, H25-Q6", "Thermal design margins below each limit (not in the RFP); apply a margin to the ~900 degC BN wall guide (current worst margin only 11.2 K)?", "", N, "PDR", "thermal"),
    ("G. H-1 head / magnet / thermal", "OQ-R8-4", "Anode thermal design target (none sourced).", "", N, "—", "anode, thermal"),

    # ---- H. Cathode C-1 (H2-2) -------------------------------------------------------------------------------------
    ("H. Cathode C-1", "HWQ-09", "Confirm Xe-fed LaB6 hollow cathode and its emission rating vs derived I_d bound + keeper + interstage currents.", "", "confirm", "—", "C-1"),
    ("H. Cathode C-1", "H22-OQ-02", "Require pulsed keeper ignition capability (300-600 V) in the supply?", "", N, "—", "PPU"),
    ("H. Cathode C-1", "H22-OQ-03", "Two isolation valves in series on the flight cathode branch?", "", N, "—", "Xe path"),
    ("H. Cathode C-1", "H22-OQ-04", "Resistor value / switching for the cathode-common tie to ground / spacecraft common.", "", N, "—", "electrical"),
    ("H. Cathode C-1", "H22-OQ-05", "Flow step and stopping rule for the spot-mode minimum-flow search.", "", "0.005 mg/s step", "—", "A7 blocker 3"),
    ("H. Cathode C-1", "H22-OQ-06", "Purge flow/duration before heating and bound on the ignition-flow dwell (unbounded analog reaches ~21 kg Xe at heavy-restart scenario).", "", N + " (bound required)", "—", "Xe mass, A7 blocker 3"),
    ("H. Cathode C-1", "H22-OQ-07", "Keeper material given O-chemistry: graphite or O-resistant alternative?", "", N, "—", "life"),
    ("H. Cathode C-1", "H22-OQ-08", "Emitter temperature floor 1843 K during O-bearing operation?", "", "adopt", "—", "C-1 life"),
    ("H. Cathode C-1", "H22-FQ-1", "Accept conservative +-2 % FS flow class (4.0 % at 0.10 mg/s, 0.216 kg Xe over 15,000 h) as the Xe-ledger flow term until S1a shows the standard class on Xe?", "", "accept", "—", "Xe mass"),
    ("H. Cathode C-1", "H22-FQ-2", "Adopt per-block installed zero check and a declared MFC body-temperature band as a test rule?", "", "adopt", "—", "metrology"),
    ("H. Cathode C-1", "H22-FQ-3", "Adopt resolution requirement <= 0.0005 mg/s and digital setpoint for the cathode MFC?", "", "adopt", "H3", "procurement"),
    ("H. Cathode C-1", "OQ-R8-5, HWQ-21", "Heated emitter witness near C-1 (metered heater load) or ground-only heated-emitter exposure?", "witness near C-1 | ground-only", N, "—", "bus power, test"),

    # ---- I. Gas path, plenum, PPU/bus (H2-3, H2-4) -----------------------------------------------------------------
    ("I. Gas path / PPU / bus", "GP-D01", "Metering-valve control strategy.", "fixed plenum setpoint | floating plenum pressure + fixed restrictor", "decide after compressor outlet characteristic known; carry both", "—", "gas path"),
    ("I. Gas path / PPU / bus", "GP-D02", "Adopt +-5 % / 10 % p-p cold-flow uniformity criterion (sets required plenum/compressor outlet pressure)?", "", "adopt for S1a", "S1a", "compressor pressure"),
    ("I. Gas path / PPU / bus", "GP-D03", "Gas-path wall recombination policy.", "catalytic path (deliver N2+O2) | inert lining (keep some atomic O)", N, "—", "feed chemistry"),
    ("I. Gas path / PPU / bus", "GP-D04", "Exclude silver from wetted parts?", "", "yes", "—", "materials"),
    ("I. Gas path / PPU / bus", "H23-Q1", "IF-A3 at the compressor outlet flange, upstream of V0 (needs upstream ICD amendment)?", "", "yes", "—", "ICD"),
    ("I. Gas path / PPU / bus", "H23-Q2", "Gas isolator withstand margin above 350 V (Paschen minimum is crossed; segmented/porous break + gas-filled withstand test proposed).", "", N, "S1a", "isolator design"),
    ("I. Gas path / PPU / bus", "HWQ-08, OQ-R8-1, OQ-R8-2, OQ-R3-3", "Anode / distributor material: accept 316L as H-1 baseline (recorded limitation) pending coupons; which candidates on coupons (316L, chromia Ni alloy, alumina-former, Rh, Pt, Cr, IrO2/RuO2, bare W control, graphite); biased (electron-collecting) or floating coupons?",
     "", "316L baseline; biased coupons", "—", "anode life (O)"),
    ("I. Gas path / PPU / bus", "OQ-R3-4, H26Q-07", "O2 cleaning standard (ASTM G93 Level C?) and O2 safety-case ownership.", "", "ASTM G93 Level C", "S1a", "safety"),
    ("I. Gas path / PPU / bus", "OQ-H24-01", "Does '< 1.5 kW' apply at bus_power_boundary_v1 and to start-up transients?", "", N, L1, "power gate"),
    ("I. Gas path / PPU / bus", "OQ-H24-02", "Must a contingency pre-ionizer fit inside the A5 power allocation, or may it use the allocation-to-requirement margin?", "inside allocation | use margin", N, L1, "A7 blocker 2"),
    ("I. Gas path / PPU / bus", "OQ-H24-03", "Adopt the proposed supply partition (discharge, per-coil magnet, cathode CS-A/CS-B, aux valve/housekeeping, reserved DC port)?", "", "adopt", "—", "PPU"),
    ("I. Gas path / PPU / bus", "OQ-H24-04", "Spacecraft bus voltage and redundancy policy.", "28 V class | regulated 100 V", N, "—", "PPU efficiency, mass"),
    ("I. Gas path / PPU / bus", "OQ-H24-05", "Sequencing SEQ-1: heater off before compressor spin-up and before any pre-ionizer seed?", "", "adopt", "—", "power peak"),
    ("I. Gas path / PPU / bus", "OQ-H24-06", "Procure a flight-representative breadboard discharge supply so eta_d is measured before LOCK-2?", "", N, "LOCK-2", "A7 blocker 2"),
    ("I. Gas path / PPU / bus", "V2-Q1", "Does the proposed 300 W common allocation include the 50 W controls/thermal?", "yes | no", N, "—", "bus power"),

    # ---- J. Diagnostics, thrust stand, facilities (H2-6, R4, R5, R7) -----------------------------------------------
    ("J. Diagnostics / stand / facility", "OD-TS-1, H26Q-02, R4-Q5", "Thrust-stand principle.", "torsional | null inverted pendulum | double/dual pendulum", N, L1, "thrust gate"),
    ("J. Diagnostics / stand / facility", "OD-TS-2", "Maximum moving mass on the stand and its spread across configurations.", "", N, L1, "stand design"),
    ("J. Diagnostics / stand / facility", "OD-TS-3", "How service lines cross the stand (liquid-metal pots, harp routing, RF/microwave hard line with shams or wireless).", "", N, "—", "stand design"),
    ("J. Diagnostics / stand / facility", "OD-TS-4", "Thrust stand: build / partner (DLR, ESA, Surrey) / buy (AST).", "build | partner | buy", N, "H3", "procurement"),
    ("J. Diagnostics / stand / facility", "OD-TS-5", "Calibration principle and force-traceability class (CSIR-NPL/NABL E2 masses vs other).", "", N, "S1a", "metrology"),
    ("J. Diagnostics / stand / facility", "OD-TS-6", "Pre-registered S1a u_T acceptance test at 12 mN with max payload and all lines installed?", "", "yes", "S1a", "thrust gate"),
    ("J. Diagnostics / stand / facility", "OD-TS-7", "Keep 1 % absolute thrust uncertainty gate (published stands show 1.4-2.2 %) or revise?", "keep | revise (value)", N, "LOCK-2", "thrust gate"),
    ("J. Diagnostics / stand / facility", "H26Q-03", "Adopt the kinematic module carrier (module weight off IP-DN, H-1 never unbolted) as the fixture concept?", "", "adopt", L1, "A7 blocker 2"),
    ("J. Diagnostics / stand / facility", "H26Q-04, R4-Q1", "MFC ranges: 4 overlapping ranges per pure-gas path (20 % FS floor) vs A4 'at least three'; does A4 authorize purchase orders or only quotations?", "", "4 ranges; quotations only until H3", "H3", "procurement"),
    ("J. Diagnostics / stand / facility", "R4-Q2", "MFC principle per gas path.", "thermal own-gas-calibrated | DP + property library | Coriolis", N, "H3", "metrology"),
    ("J. Diagnostics / stand / facility", "R4-Q3", "Cathode-Xe MFC: cover only 0.05-0.2 mg/s, or also start/diode flows (0.6-0.8 mg/s) - one or two controllers?", "one | two", N, "H3", "procurement"),
    ("J. Diagnostics / stand / facility", "R4-Q4", "In-house rate-of-rise calibrator acceptable as S1a primary flow standard, or NABL-certified MFC calibration required?", "", N, "S1a", "metrology"),
    ("J. Diagnostics / stand / facility", "R4-Q6", "RGA in the S1a minimum set? If yes, 200 or 300 amu, differential pumping?", "", N, "S1a", "instrumentation"),
    ("J. Diagnostics / stand / facility", "WEB-INS-1", "C-1 pyrometer view (HW-C1-09(b)).", "", N, "—", "instrumentation"),
    ("J. Diagnostics / stand / facility", "H26Q-05", "Exploratory I_d(t) chain to 60 MHz for S1b, or a narrower band recorded as such?", "", N, "S1b", "oscillation data"),
    ("J. Diagnostics / stand / facility", "H26Q-06", "Flight telemetry subset (ground vs flight diagnostics) as input to PPU / control-FDIR.", "", "accept H2-6 list", "—", "FDIR"),
    ("J. Diagnostics / stand / facility", "H26Q-08", "Ground radiative sink temperature for the thermal model (facility cryopanel/wall).", "", N, "—", "thermal"),
    ("J. Diagnostics / stand / facility", "HWQ-11", "Atomic-O representativeness of Phase 3: accept N2 + O2 surrogate with recorded limitation, or add a dedicated atomic-O source?", "surrogate | add AO source", N, "—", "Phase 3"),
    ("J. Diagnostics / stand / facility", "HWQ-12", "Sham service lines on the stand in every configuration?", "", "yes", "—", "A7 blocker 2"),
    ("J. Diagnostics / stand / facility", "HWQ-16", "Witness coupons/holders = non-functional exchangeable items (no new H-1') if B(z) and HW-0 reference unchanged?", "", "yes", "—", "test validity"),
    ("J. Diagnostics / stand / facility", "HWQ-17", "Approve alternative-grade wall sector inserts (N/O wall data, but changes design-representativeness)?", "", N, "—", "H-1 representativeness"),
    ("J. Diagnostics / stand / facility", "R5-Q1", "Foreign facility (e.g. SITAEL IV10) admissible for score-bearing data (TDF rules, export of H-1, data custody, competitors)?", "", N, L1, "facility choice"),
    ("J. Diagnostics / stand / facility", "R5-Q2", "Request ISRO LPSC facility specifications through official channels before T-PB-MAX is fixed?", "", N, L1, "facility choice"),
    ("J. Diagnostics / stand / facility", "R5-Q5", "XPS / light-element EDS: accredited lab mandatory, or academic lab with results labelled semi-quantitative?", "", N, "—", "analysis"),
    ("J. Diagnostics / stand / facility", "R5-Q6", "Engineering-only S1a in a smaller domestic chamber, score-bearing stages at one facility meeting REQ-FAC-05?", "", N, "S1a", "facility plan"),

    # ---- K. Governance: M16 / register ----------------------------------------------------------------------------
    ("K. Governance", "M16-Q-01", "Assign an accountable owner to each of the 17 subsystem rows.", "", N, "—", "none"),
    ("K. Governance", "M16-Q-02, REFRESH-Q1", "Accept the scheduler rule (READY/RUNNING/BLOCKED/VERIFIED, one blocking item per row, non-lane-input guard), per-row blocking items and A7 categories?", "", "accept", "—", "none"),
    ("K. Governance", "M16-Q-03", "Accept the H2-lane-to-row mapping?", "", "accept", "—", "none"),
    ("K. Governance", "M16-Q-04", "Direct owning lanes to resolve reconciliation items RC-01..RC-10 (or accept divergences).", "", "direct", "—", "see items"),
    ("K. Governance", "REFRESH-Q3", "Promote inferred gates (before S1 / HRR / Phase 1) to stated deadlines by mapping them to LOCK-1/LOCK-2?", "", N, "—", "none"),

    # ---- L. Parallel RF||Hall v2 ---------------------------------------------------------------------------------
    ("L. Parallel RF||Hall v2", "V2-Q2", "Supply the v2 inputs listed in VALIDATION_PLAN.md (RF coupling, nozzle, Hall-Xe points, compressor, start-up Xe).", "", N, "—", "v2 surfaces"),
    ("L. Parallel RF||Hall v2", "V2-Q3", "Generate hypothesis-grade sec. 45 surfaces from inputs you label ASSUMED_SCREENING_VALUE?", "yes | no", "no until inputs exist", "—", "v2"),
    ("L. Parallel RF||Hall v2", "V2-Q4", "A8 (parallel RF||Hall investigation) stays DRAFT_ARCHITECTURE_HYPOTHESIS, or freeze as an owner decision?", "keep draft | freeze", N, "—", "architecture docs"),
]

HEAD = ["No", "Area", "Covers IDs", "Question", "Options", "Proposed", "Needed by", "Affects", "Your answer"]


def table():
    return [[i + 1, *r, ""] for i, r in enumerate(ROWS)]


def write_md(rows, path):
    def esc(s):
        return str(s).replace("|", "\\|").replace("\n", " ")
    lines = ["# Consolidated owner questions", "",
             f"{len(rows)} decisions (duplicates across sources merged; 'Covers IDs' lists every source id a row answers).",
             "Generated by build_owner_questions_consolidated.py. 'Proposed' is the source's proposal, not a decision.",
             "Answer by row number or ID, e.g. '12 accept, 30 B, 55 no'.", "",
             "| " + " | ".join(HEAD) + " |", "|" + "---|" * len(HEAD)]
    lines += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows]
    open(path, "w", encoding="utf-8").write("\n".join(lines) + "\n")


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(HEAD)
        w.writerows(rows)


def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook()
    ws = wb.active
    ws.title = "Owner questions"
    f = Font(name="Arial", size=10)
    fb = Font(name="Arial", size=10, bold=True)
    ws.append(["Owner questions - fill in the yellow 'Your answer' column (column I) only. 'Proposed' is the source's "
               "proposal, not a decision. Answers like 'accept', an option letter, or free text are fine."])
    ws["A1"].font = fb
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    widths = [5, 22, 24, 70, 45, 38, 14, 24, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    yellow = PatternFill("solid", start_color="FFFF00")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = fb if c.row == 2 else f
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[0].row > 2:
            row[8].fill = yellow
    ws.freeze_panes = "E3"
    t = Table(displayName="OwnerQuestions", ref=f"A2:I{ws.max_row}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    wb.save(path)


if __name__ == "__main__":
    rows = table()
    write_md(rows, os.path.join(HERE, "OWNER_QUESTIONS_CONSOLIDATED.md"))
    write_csv(rows, os.path.join(HERE, "owner_questions_consolidated.csv"))
    write_xlsx(rows, os.path.join(HERE, "owner_questions_consolidated.xlsx"))
    print(f"{len(rows)} rows written")
